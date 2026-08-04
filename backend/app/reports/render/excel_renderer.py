"""Excel renderer from report sections (openpyxl)."""

from __future__ import annotations

import io
from typing import Any


def render_excel(title: str, sections: list[dict[str, Any]]) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    if not sections:
        ws = wb.create_sheet("Report")
        ws.append([title])
    for idx, section in enumerate(sections):
        name = str(section.get("metricKey") or f"Sheet{idx + 1}")[:31]
        ws = wb.create_sheet(name)
        cols = section.get("columns") or []
        rows = section.get("rows") or []
        if cols:
            ws.append(cols)
        for row in rows:
            ws.append([str(c) for c in row])
        if not cols and not rows:
            ws.append([title, "(no data)"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
